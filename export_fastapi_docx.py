#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
将 FastAPI 的 OpenAPI 接口定义导出为 Word（.docx）接口说明书。

依赖：
    pip install python-docx

使用：
    python export_fastapi_docx.py --app main:app
    python export_fastapi_docx.py --app app.main:app --output docs/API接口文档.docx
    python export_fastapi_docx.py --url http://127.0.0.1:8000/openapi.json
    python export_fastapi_docx.py --input openapi.json

注意：--app 模式请从项目根目录运行，确保能够导入你的 FastAPI 应用。
"""
import argparse
import importlib
import json
import sys
from collections import defaultdict
from datetime import datetime
from pathlib import Path
from urllib.request import Request, urlopen

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


HTTP_METHODS = ("get", "post", "put", "patch", "delete", "options", "head", "trace")
MAX_SCHEMA_DEPTH = 6
CHINESE_FONT = "微软雅黑"
TEXT_FONT = "Calibri"


def set_text_fonts(r_pr, latin_font=TEXT_FONT):
    """显式设置中西文字体，避免模板的主题字体覆盖中文字体。"""
    fonts = r_pr.get_or_add_rFonts()
    for attribute in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        fonts.attrib.pop(qn("w:" + attribute), None)
    for attribute in ("ascii", "hAnsi", "cs"):
        fonts.set(qn("w:" + attribute), latin_font)
    fonts.set(qn("w:eastAsia"), CHINESE_FONT)
    fonts.set(qn("w:hint"), "eastAsia")
    language = r_pr.find(qn("w:lang"))
    if language is None:
        language = OxmlElement("w:lang")
        r_pr.append(language)
    language.set(qn("w:val"), "en-US")
    language.set(qn("w:eastAsia"), "zh-CN")


def as_text(value):
    if value is None:
        return ""
    if isinstance(value, str):
        return value
    return json.dumps(value, ensure_ascii=False)


def ref_value(spec, ref):
    """解析 #/components/schemas/Model 等 OpenAPI 内部引用。"""
    if not isinstance(ref, str) or not ref.startswith("#/"):
        return {}
    node = spec
    for part in ref[2:].split("/"):
        part = part.replace("~1", "/").replace("~0", "~")
        if not isinstance(node, dict) or part not in node:
            return {}
        node = node[part]
    return node if isinstance(node, dict) else {}


def resolve(spec, value):
    """展开直接 $ref，保留相邻的 description 等字段。"""
    if not isinstance(value, dict):
        return {}
    result = value
    seen = set()
    while "$ref" in result:
        ref = result["$ref"]
        if ref in seen:
            break
        seen.add(ref)
        target = ref_value(spec, ref)
        if not target:
            break
        result = {**target, **{k: v for k, v in result.items() if k != "$ref"}}
    return result


def schema_type(schema):
    """生成适合展示的类型名称，不需要展开引用。"""
    if not isinstance(schema, dict):
        return "未声明"
    if "$ref" in schema:
        return schema["$ref"].rsplit("/", 1)[-1]
    for key in ("anyOf", "oneOf", "allOf"):
        if key in schema:
            return " | ".join(schema_type(s) for s in schema[key])
    kind = schema.get("type")
    if isinstance(kind, list):
        return " | ".join(str(x) for x in kind)
    if kind == "array":
        return f"array[{schema_type(schema.get('items', {}))}]"
    if kind == "object":
        extra = schema.get("additionalProperties")
        return f"map[string, {schema_type(extra)}]" if isinstance(extra, dict) else "object"
    if not kind and "properties" in schema:
        return "object"
    fmt = schema.get("format")
    return f"{kind} ({fmt})" if kind and fmt else (kind or "未声明")


def field_rows(spec, schema, prefix="", visited=None, depth=0):
    """将嵌套的 Pydantic/OpenAPI 模型展开为字段表；限制递归深度。"""
    if not isinstance(schema, dict) or depth >= MAX_SCHEMA_DEPTH:
        return []
    visited = visited or frozenset()
    ref = schema.get("$ref")
    if ref in visited:
        return []
    visited = visited | ({ref} if ref else set())
    value = resolve(spec, schema)
    rows = []
    required = set(value.get("required", []))

    for name, field in value.get("properties", {}).items():
        if not isinstance(field, dict):
            continue
        label = f"{prefix}.{name}" if prefix else name
        actual = resolve(spec, field)
        note = field.get("description") or actual.get("description") or ""
        if "enum" in actual:
            note += f"；可选值：{as_text(actual['enum'])}"
        if "default" in actual:
            note += f"；默认值：{as_text(actual['default'])}"
        rows.append([
            label, schema_type(field),
            "是" if name in required else "否", note.strip("；")
        ])
        child = actual.get("items") if actual.get("type") == "array" else field
        next_prefix = label + "[]" if actual.get("type") == "array" else label
        rows.extend(field_rows(spec, child, next_prefix, visited, depth + 1))

    for key in ("allOf", "anyOf", "oneOf"):
        for item in value.get(key, []):
            if item.get("type") == "null":
                continue
            rows.extend(field_rows(spec, item, prefix, visited, depth + 1))

    if value.get("type") == "array" and not value.get("properties"):
        rows.extend(field_rows(
            spec, value.get("items", {}), prefix or "items[]", visited, depth + 1
        ))

    # allOf/anyOf 可能引入重复字段，按字段路径去重。
    unique = {}
    for row in rows:
        unique.setdefault(row[0], row)
    return list(unique.values())


def sample_from_schema(spec, schema, visited=None, depth=0):
    """按类型合成结构示例；合成值不代表真实业务返回数据。"""
    if not isinstance(schema, dict) or depth >= MAX_SCHEMA_DEPTH:
        return None
    visited = visited or frozenset()
    ref = schema.get("$ref")
    if ref in visited:
        return "<递归引用>"
    visited = visited | ({ref} if ref else set())
    value = resolve(spec, schema)

    for key in ("example", "default", "const"):
        if key in value:
            return value[key]
    if value.get("examples") and isinstance(value["examples"], list):
        return value["examples"][0]
    if value.get("enum"):
        return value["enum"][0]

    if "allOf" in value:
        samples = [sample_from_schema(spec, x, visited, depth + 1) for x in value["allOf"]]
        if all(isinstance(x, dict) for x in samples):
            merged = {}
            for x in samples:
                merged.update(x)
            return merged
        return next((x for x in samples if x is not None), None)

    for key in ("anyOf", "oneOf"):
        if key in value:
            choices = [x for x in value[key] if x.get("type") != "null"]
            return sample_from_schema(spec, choices[0], visited, depth + 1) if choices else None

    kind = value.get("type", "object" if "properties" in value else None)
    if isinstance(kind, list):
        kind = next((x for x in kind if x != "null"), "null")
    if kind == "object":
        props = value.get("properties", {})
        if props:
            return {k: sample_from_schema(spec, v, visited, depth + 1)
                    for k, v in props.items()}
        extra = value.get("additionalProperties")
        return {"key": sample_from_schema(spec, extra, visited, depth + 1)} \
            if isinstance(extra, dict) else {}
    if kind == "array":
        item = sample_from_schema(spec, value.get("items", {}), visited, depth + 1)
        return [item] if item is not None else []
    if kind == "string":
        fmt = value.get("format")
        return {"date": "2026-01-01", "date-time": "2026-01-01T00:00:00Z",
                "email": "user@example.com"}.get(fmt, "string")
    if kind == "integer":
        return 0
    if kind == "number":
        return 0.0
    if kind == "boolean":
        return False
    return None


def media_example(spec, content):
    """优先使用开发者定义的 example / examples，否则生成结构示例。"""
    if "example" in content:
        return content["example"], "示例（接口定义提供）"
    examples = content.get("examples", {})
    if isinstance(examples, dict):
        for val in examples.values():
            item = resolve(spec, val)
            if "value" in item:
                return item["value"], "示例（接口定义提供）"
    schema = content.get("schema", {})
    expanded = resolve(spec, schema)
    if "example" in expanded:
        return expanded["example"], "示例（接口定义提供）"
    if isinstance(expanded.get("examples"), list) and expanded["examples"]:
        return expanded["examples"][0], "示例（接口定义提供）"
    return sample_from_schema(spec, schema), "结构示例（自动推断，非实际响应）"


def set_cell_background(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    element = OxmlElement("w:tblHeader")
    element.set(qn("w:val"), "true")
    tr_pr.append(element)


def add_table(doc, headers, rows):
    if not rows:
        doc.add_paragraph("无")
        return
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = True
    set_repeat_header(table.rows[0])
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = header
        set_cell_background(cell, "E7EEF6")
        for p in cell.paragraphs:
            for run in p.runs:
                run.bold = True
                run.font.size = Pt(9)
                set_text_fonts(run._element.get_or_add_rPr())
    for item in rows:
        cells = table.add_row().cells
        for i, content in enumerate(item):
            cells[i].text = as_text(content)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cells[i].paragraphs:
                p.paragraph_format.line_spacing = 1.15
                for run in p.runs:
                    run.font.size = Pt(9)
                    set_text_fonts(run._element.get_or_add_rPr())
    doc.add_paragraph()  # 表格与下一段之间留出间距


def add_code_example(doc, spec, media_type, content):
    if "json" not in media_type.lower():
        return
    value, title = media_example(spec, content)
    if value is None:
        return
    doc.add_paragraph(title + "：")
    code = json.dumps(value, ensure_ascii=False, indent=2)
    if len(code) > 4500:
        code = code[:4500] + "\n...（示例过长，已截断）"
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Cm(0.3)
    paragraph.paragraph_format.right_indent = Cm(0.3)
    paragraph.paragraph_format.line_spacing = 1.15
    run = paragraph.add_run(code)
    set_text_fonts(run._element.get_or_add_rPr(), latin_font="Consolas")
    run.font.size = Pt(8)


def setup_document(doc):
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)  # A4
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin = section.right_margin = Cm(1.8)

    # 默认字体同时覆盖正文、表格和页脚中未直接设置字体的文字。
    defaults = doc.styles.element.find(qn("w:docDefaults"))
    if defaults is None:
        defaults = OxmlElement("w:docDefaults")
        doc.styles.element.insert(0, defaults)
    run_defaults = defaults.find(qn("w:rPrDefault"))
    if run_defaults is None:
        run_defaults = OxmlElement("w:rPrDefault")
        defaults.insert(0, run_defaults)
    r_pr = run_defaults.find(qn("w:rPr"))
    if r_pr is None:
        r_pr = OxmlElement("w:rPr")
        run_defaults.append(r_pr)
    set_text_fonts(r_pr)

    for name, size in (("Normal", 10), ("Title", 20), ("Heading 1", 15),
                       ("Heading 2", 12), ("Heading 3", 10.5)):
        style = doc.styles[name]
        set_text_fonts(style._element.get_or_add_rPr())
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(25, 49, 78) if name != "Normal" else RGBColor(38, 47, 59)
        style.paragraph_format.line_spacing = 1.25

    for name in ("Default Paragraph Font", "Normal Table", "Table Grid", "Footer",
                 "Footer Char", "Title Char", "Heading 1 Char", "Heading 2 Char",
                 "Heading 3 Char"):
        if name in doc.styles:
            set_text_fonts(doc.styles[name]._element.get_or_add_rPr())

    footer = section.footer.paragraphs[0]
    footer.alignment = 1
    footer.add_run("由 FastAPI OpenAPI 自动生成 · " + datetime.now().strftime("%Y-%m-%d"))


def get_parameters(spec, path_item, operation):
    merged = {}
    for raw in path_item.get("parameters", []) + operation.get("parameters", []):
        param = resolve(spec, raw)
        merged[(param.get("name"), param.get("in"))] = param
    rows = []
    for param in merged.values():
        schema = param.get("schema", {})
        note = param.get("description", "")
        actual = resolve(spec, schema)
        if "enum" in actual:
            note += f"；可选值：{as_text(actual['enum'])}"
        if "default" in actual:
            note += f"；默认值：{as_text(actual['default'])}"
        rows.append([
            param.get("name", ""), param.get("in", ""),
            schema_type(schema), "是" if param.get("required") else "否",
            note.strip("；"),
        ])
    return rows


def get_security(spec, operation):
    security = operation.get("security", spec.get("security", []))
    if not security:
        return "未要求（以接口定义为准）"
    return " 或 ".join(
        " + ".join(schemes.keys()) if schemes else "无需认证"
        for schemes in security
    )


def write_operation(doc, spec, path, method, path_item, operation, index):
    label = operation.get("summary") or operation.get("operationId") or "未命名接口"
    if operation.get("deprecated"):
        label += "（已废弃）"
    doc.add_heading(f"{index}. {label}", level=2)

    add_table(doc, ["项目", "内容"], [
        ["请求方法", method.upper()],
        ["接口地址", path],
        ["接口分组", "、".join(operation.get("tags", ["未分组"]))],
        ["鉴权方式", get_security(spec, operation)],
    ])

    if operation.get("description"):
        doc.add_paragraph(operation["description"])

    doc.add_heading("请求参数（路径 / 查询 / Header / Cookie）", level=3)
    add_table(doc, ["参数名", "位置", "类型", "必填", "说明"],
              get_parameters(spec, path_item, operation))

    doc.add_heading("请求体", level=3)
    body = resolve(spec, operation.get("requestBody", {}))
    content_map = body.get("content", {})
    if not content_map:
        doc.add_paragraph("无")
    else:
        doc.add_paragraph("是否必填：" + ("是" if body.get("required") else "否"))
        if body.get("description"):
            doc.add_paragraph(body["description"])
        for media_type, content in content_map.items():
            doc.add_paragraph(f"Content-Type：{media_type}")
            schema = content.get("schema", {})
            rows = field_rows(spec, schema)
            if not rows:
                rows = [["(整个请求体)", schema_type(schema), "—", ""]]
            add_table(doc, ["字段名", "类型", "必填", "说明"], rows)
            add_code_example(doc, spec, media_type, content)

    doc.add_heading("响应状态", level=3)
    response_rows = []
    for status, raw in operation.get("responses", {}).items():
        response = resolve(spec, raw)
        descriptions = []
        for media_type, content in response.get("content", {}).items():
            descriptions.append(f"{media_type}: {schema_type(content.get('schema', {}))}")
        response_rows.append([status, response.get("description", ""),
                              "；".join(descriptions) or "无响应体"])
    add_table(doc, ["状态码", "说明", "响应类型"], response_rows)

    for status, raw in operation.get("responses", {}).items():
        response = resolve(spec, raw)
        for media_type, content in response.get("content", {}).items():
            doc.add_heading(f"响应 {status}（{media_type}）", level=3)
            schema = content.get("schema", {})
            rows = field_rows(spec, schema)
            if not rows:
                rows = [["(整个响应体)", schema_type(schema), "—", ""]]
            add_table(doc, ["字段名", "类型", "必填", "说明"], rows)
            add_code_example(doc, spec, media_type, content)


def export_docx(spec, output):
    if not isinstance(spec, dict) or not isinstance(spec.get("paths"), dict):
        raise ValueError("输入内容不是有效的 OpenAPI 文档（缺少 paths）。")

    doc = Document()
    setup_document(doc)
    info = spec.get("info", {})
    title = info.get("title", "FastAPI 接口文档")
    doc.core_properties.title = title + " — API 接口说明书"

    doc.add_heading(title + " — API 接口说明书", level=0)
    doc.add_paragraph("API 版本：" + str(info.get("version", "未声明")))
    doc.add_paragraph("文档生成日期：" + datetime.now().strftime("%Y-%m-%d %H:%M"))
    if info.get("description"):
        doc.add_paragraph(info["description"])
    if spec.get("servers"):
        doc.add_paragraph("服务地址：" + "、".join(
            x.get("url", "") for x in spec["servers"]
        ))

    groups = defaultdict(list)
    for path, path_item in spec["paths"].items():
        if not isinstance(path_item, dict):
            continue
        for method in HTTP_METHODS:
            operation = path_item.get(method)
            if isinstance(operation, dict):
                group = (operation.get("tags") or ["未分组"])[0]
                groups[group].append((path, method, path_item, operation))

    doc.add_paragraph(f"共 {sum(map(len, groups.values()))} 个接口，{len(groups)} 个接口分组。")
    tag_descriptions = {
        t.get("name"): t.get("description", "")
        for t in spec.get("tags", [])
    }

    for group_name, operations in groups.items():
        doc.add_heading(group_name, level=1)
        if tag_descriptions.get(group_name):
            doc.add_paragraph(tag_descriptions[group_name])
        for number, (path, method, path_item, operation) in enumerate(operations, 1):
            write_operation(doc, spec, path, method, path_item, operation, number)

    if not groups:
        doc.add_paragraph("未发现可导出的接口；请检查 include_in_schema 和应用的路由注册情况。")

    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output))
    return output.resolve(), sum(map(len, groups.values()))


def load_openapi(args):
    if args.input:
        with open(args.input, "r", encoding="utf-8-sig") as file:
            return json.load(file)
    if args.url:
        request = Request(args.url, headers={"Accept": "application/json"})
        with urlopen(request, timeout=20) as response:
            return json.load(response)

    # 即使脚本放在项目目录之外，也优先从当前工作目录导入 FastAPI 项目。
    cwd = str(Path.cwd())
    if cwd not in sys.path:
        sys.path.insert(0, cwd)
    target = args.app or "main:app"
    if ":" not in target:
        raise ValueError("--app 格式应为 module:app，例如 main:app 或 app.main:app")
    module_name, app_name = target.split(":", 1)
    module = importlib.import_module(module_name)
    app = getattr(module, app_name)
    if not callable(getattr(app, "openapi", None)):
        raise TypeError(f"{target} 不是有效的 FastAPI 应用（缺少 .openapi()）。")
    return app.openapi()


def main():
    parser = argparse.ArgumentParser(description="FastAPI OpenAPI → Word API 接口说明书")
    group = parser.add_mutually_exclusive_group()
    group.add_argument("--app", help="应用模块:对象名，例如 main:app（默认）")
    group.add_argument("--url", help="OpenAPI JSON 地址，例如 http://127.0.0.1:8000/openapi.json")
    group.add_argument("--input", help="本地 openapi.json 文件路径")
    parser.add_argument("--output", default="docs/API接口文档.docx", help="导出的 docx 文件名")
    args = parser.parse_args()

    try:
        openapi = load_openapi(args)
        output, count = export_docx(openapi, args.output)
        print(f"✅ 导出成功：{output}")
        print(f"   共导出 {count} 个接口")
    except Exception as exc:
        print(f"❌ 生成失败：{type(exc).__name__}: {exc}", file=sys.stderr)
        print("提示：--app 模式须在项目根目录执行；也可以改用 --url。", file=sys.stderr)
        sys.exit(1)


if __name__ == "__main__":
    main()
